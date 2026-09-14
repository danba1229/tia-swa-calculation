from __future__ import annotations

import hashlib
import io
import json
import math
import os
import struct
import zipfile
from array import array
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:  # Runtime calculation only reads verified derived files.
    load_workbook = None


ROOT = Path(os.getenv("TIA_COMMON_DATA_ROOT") or Path(__file__).resolve().parents[1]).resolve()
COMMON = ROOT / "common_data"
RAW = COMMON / "raw"
DERIVED = COMMON / "derived"
ZONE_WORKBOOK = ROOT / "downloads" / "metro_2024_zone_system_verified.xlsx"
BOUNDARY_GEOJSON = RAW / "HangJeongDong_ver20230101.geojson"
BOUNDARY_REFERENCE_GEOJSON = RAW / "HangJeongDong_ver20220101.geojson"
BOUNDARY_FALLBACK_GEOJSON = BOUNDARY_REFERENCE_GEOJSON
PURPOSE_OD = ROOT / "downloads" / "ODTRIP23_F_verified.txt"
ZONE_JSON = DERIVED / "metro_2024_1310_zones.json"
MATRIX_BIN = DERIVED / "metro_2024_1310_purpose_total_2023.f64"
PURPOSE_OUTBOUND_JSON = DERIVED / "metro_2024_1310_purpose_outbound_2023.json"
NATIONAL_INDICATOR_JSON = DERIVED / "national_2024_250_indicators.json"
NATIONAL_ZONE_JSON = DERIVED / "national_2024_250_zones.json"
REGION_CROSSWALK_JSON = DERIVED / "metro_2024_to_national_2024_region_crosswalk.json"
MANIFEST_JSON = DERIVED / "manifest.json"
ZONE_COUNT = 1310
PURPOSES = ("귀가", "출근", "등교", "업무", "기타")
NATIONAL_YEARS = (2023, 2025, 2030, 2035, 2040, 2045, 2050)
NATIONAL_PURPOSE_BUNDLE = ROOT.parent / "_analysis_readonly" / "source" / "2024-OD-PSN-OBJ-00 전국지역간 목적 OD(250존)(2023-2050).zip"
NATIONAL_MODE_BUNDLE = ROOT.parent / "_analysis_readonly" / "source" / "2024-OD-PSN-MOD-10 전국지역간 주수단 OD(250존)(2023-2050).zip"


def national_matrix_path(year: int) -> Path:
    return DERIVED / f"national_2024_250_purpose_total_{int(year)}.f64"


class CommonDataError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _ring_measure(ring: list[list[float]]) -> tuple[float, float, float]:
    """Return absolute planar area and centroid for a WGS84 ring."""
    signed_twice_area = 0.0
    cx_acc = 0.0
    cy_acc = 0.0
    for first, second in zip(ring, ring[1:] + ring[:1]):
        cross = float(first[0]) * float(second[1]) - float(second[0]) * float(first[1])
        signed_twice_area += cross
        cx_acc += (float(first[0]) + float(second[0])) * cross
        cy_acc += (float(first[1]) + float(second[1])) * cross
    if abs(signed_twice_area) < 1e-18:
        xs = [float(point[0]) for point in ring]
        ys = [float(point[1]) for point in ring]
        return 0.0, math.fsum(xs) / len(xs), math.fsum(ys) / len(ys)
    return abs(signed_twice_area) / 2.0, cx_acc / (3.0 * signed_twice_area), cy_acc / (3.0 * signed_twice_area)


def _geometry_measure(geometry: dict[str, Any]) -> tuple[float, float, float]:
    polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
    weighted_x: list[float] = []
    weighted_y: list[float] = []
    weights: list[float] = []
    for polygon in polygons:
        if not polygon:
            continue
        outer_area, outer_x, outer_y = _ring_measure(polygon[0])
        polygon_area = outer_area
        polygon_x = outer_x * outer_area
        polygon_y = outer_y * outer_area
        for hole in polygon[1:]:
            hole_area, hole_x, hole_y = _ring_measure(hole)
            polygon_area -= hole_area
            polygon_x -= hole_x * hole_area
            polygon_y -= hole_y * hole_area
        if polygon_area > 0:
            weights.append(polygon_area)
            weighted_x.append(polygon_x)
            weighted_y.append(polygon_y)
    total = math.fsum(weights)
    if total <= 0:
        raise CommonDataError("Boundary geometry has no positive polygon area")
    return total, math.fsum(weighted_x) / total, math.fsum(weighted_y) / total


def _point_in_ring(lon: float, lat: float, ring: list[list[float]]) -> bool:
    inside = False
    j = len(ring) - 1
    for i, point in enumerate(ring):
        xi, yi = float(point[0]), float(point[1])
        xj, yj = float(ring[j][0]), float(ring[j][1])
        if ((yi > lat) != (yj > lat)) and lon < (xj - xi) * (lat - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def point_in_geometry(lon: float, lat: float, geometry: dict[str, Any]) -> bool:
    polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
    for polygon in polygons:
        if polygon and _point_in_ring(lon, lat, polygon[0]):
            if not any(_point_in_ring(lon, lat, hole) for hole in polygon[1:]):
                return True
    return False


def point_on_geometry_boundary(lon: float, lat: float, geometry: dict[str, Any], tolerance: float = 1e-9) -> bool:
    polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
    for polygon in polygons:
        for ring in polygon:
            for first, second in zip(ring, ring[1:] + ring[:1]):
                x1, y1 = float(first[0]), float(first[1])
                x2, y2 = float(second[0]), float(second[1])
                dx, dy = x2 - x1, y2 - y1
                if dx == 0 and dy == 0:
                    distance = math.hypot(lon - x1, lat - y1)
                else:
                    ratio = max(0.0, min(1.0, ((lon - x1) * dx + (lat - y1) * dy) / (dx * dx + dy * dy)))
                    distance = math.hypot(lon - (x1 + ratio * dx), lat - (y1 + ratio * dy))
                if distance <= tolerance:
                    return True
    return False


def load_zone_workbook(path: Path = ZONE_WORKBOOK) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = tuple(next(rows)[:6])
    expected = ("시도", "시군구", "행정동", "권역 존체계_읍면동", "행정기관코드_읍면동", "권역내부=1, 권역외부=2")
    if header != expected:
        workbook.close()
        raise CommonDataError(f"Unexpected metro zone header: {header!r}")
    result = []
    for values in rows:
        result.append({
            "province": values[0],
            "city_county": values[1],
            "admin_dong": values[2],
            "taz": int(values[3]),
            "admin_code": str(int(values[4])),
            "scope_flag": int(values[5]),
        })
    workbook.close()
    ids = [row["taz"] for row in result]
    if len(result) != ZONE_COUNT or ids != list(range(1, ZONE_COUNT + 1)):
        raise CommonDataError("Metro zone workbook is not exactly sequential TAZ 1..1310")
    return result


def _representative_from_features(features: Iterable[dict[str, Any]]) -> tuple[float | None, float | None, str]:
    measured = []
    for feature in features:
        try:
            measured.append((*_geometry_measure(feature["geometry"]), feature["geometry"]))
        except (KeyError, TypeError, CommonDataError):
            continue
    if not measured:
        return None, None, "NO_MATCHING_BOUNDARY"
    total = math.fsum(item[0] for item in measured)
    lon = math.fsum(item[0] * item[1] for item in measured) / total
    lat = math.fsum(item[0] * item[2] for item in measured) / total
    if any(point_in_geometry(lon, lat, item[3]) for item in measured):
        return lon, lat, "AREA_WEIGHTED_BOUNDARY_CENTROID"
    largest = max(measured, key=lambda item: item[0])
    if point_in_geometry(largest[1], largest[2], largest[3]):
        return largest[1], largest[2], "LARGEST_PART_CENTROID_FALLBACK"
    geometry = largest[3]
    polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
    largest_ring = max((polygon[0] for polygon in polygons if polygon), key=len)
    for fraction in (0.5, 0.4, 0.6, 0.3, 0.7):
        xs = [float(point[0]) for point in largest_ring]
        ys = [float(point[1]) for point in largest_ring]
        candidate = (min(xs) + (max(xs) - min(xs)) * fraction, min(ys) + (max(ys) - min(ys)) * fraction)
        if point_in_geometry(candidate[0], candidate[1], geometry):
            return candidate[0], candidate[1], "INTERIOR_GRID_FALLBACK"
    return float(largest_ring[0][0]), float(largest_ring[0][1]), "BOUNDARY_VERTEX_FALLBACK"


def build_zone_points(zone_path: Path = ZONE_WORKBOOK, boundary_path: Path = BOUNDARY_GEOJSON) -> dict[str, Any]:
    zones = load_zone_workbook(zone_path)
    boundaries = json.loads(boundary_path.read_text(encoding="utf-8"))["features"]
    by_admin: dict[str, list[dict[str, Any]]] = {}
    by_sgg: dict[str, list[dict[str, Any]]] = {}
    for feature in boundaries:
        properties = feature.get("properties") or {}
        admin_code = str(properties.get("adm_cd2") or "")
        sgg_code = str(properties.get("sgg") or "")
        if admin_code:
            by_admin.setdefault(admin_code, []).append(feature)
        if sgg_code:
            by_sgg.setdefault(sgg_code, []).append(feature)
    fallback_boundaries = json.loads(BOUNDARY_FALLBACK_GEOJSON.read_text(encoding="utf-8"))["features"]
    fallback_by_admin: dict[str, list[dict[str, Any]]] = {}
    for feature in fallback_boundaries:
        admin_code = str((feature.get("properties") or {}).get("adm_cd2") or "")
        if admin_code:
            fallback_by_admin.setdefault(admin_code, []).append(feature)
    missing = []
    methods: dict[str, int] = {}
    for zone in zones:
        matches = by_admin.get(zone["admin_code"], []) if zone["scope_flag"] == 1 else by_sgg.get(zone["admin_code"], [])
        fallback_used = False
        if zone["scope_flag"] == 1 and not matches:
            matches = fallback_by_admin.get(zone["admin_code"], [])
            fallback_used = bool(matches)
        lon, lat, method = _representative_from_features(matches)
        if fallback_used:
            method = f"20220101_REFERENCE_{method}"
        zone["longitude"] = lon
        zone["latitude"] = lat
        zone["coordinate_method"] = method
        zone["boundary_version"] = "20220101" if fallback_used else "20230101"
        zone["region_name"] = " ".join(str(value) for value in (zone["province"], zone["city_county"], zone["admin_dong"]) if value)
        methods[method] = methods.get(method, 0) + 1
        if lon is None or lat is None:
            missing.append(zone["taz"])
    return {
        "schema_version": 1,
        "zone_system": "metro-2024-1310",
        "zone_count": ZONE_COUNT,
        "coordinate_source": {
            "primary_file": str(boundary_path.relative_to(ROOT)).replace("\\", "/"),
            "primary_sha256": sha256_file(boundary_path),
            "primary_version": "20230101",
            "fallback_file": str(BOUNDARY_FALLBACK_GEOJSON.relative_to(ROOT)).replace("\\", "/"),
            "fallback_sha256": sha256_file(BOUNDARY_FALLBACK_GEOJSON),
            "fallback_version": "20220101",
            "provenance": "SGIS/행정안전부 기반 vuski/admdongkor 가공본",
            "license": "CC BY 4.0 / 공공누리 제1유형 출처표시",
            "warning": "공식 배포 원본 자체가 아닌 SGIS 기반 가공 경계이며, 2023 경계에 없는 구 코드 6개만 2022-01-01 기준 경계로 보완",
        },
        "representative_point_policy": "polygon area-weighted centroid; fall back to largest part/interior point when needed",
        "coordinate_method_counts": methods,
        "missing_coordinate_taz": missing,
        "zones": zones,
    }


def build_purpose_matrix(od_path: Path = PURPOSE_OD) -> dict[str, Any]:
    matrix = array("d", [0.0]) * (ZONE_COUNT * ZONE_COUNT)
    seen = bytearray(ZONE_COUNT * ZONE_COUNT)
    outbound = [[0.0] * len(PURPOSES) for _ in range(ZONE_COUNT)]
    origin_zone_ids: dict[int, str] = {}
    destination_zone_ids: dict[int, str] = {}
    record_count = 0
    with od_path.open("r", encoding="cp949") as stream:
        for line_number, line in enumerate(stream, start=1):
            fields = line.split()
            if len(fields) != 9:
                raise CommonDataError(f"OD line {line_number} has {len(fields)} fields, expected 9")
            origin, destination = int(fields[0]), int(fields[2])
            if not 1 <= origin <= ZONE_COUNT or not 1 <= destination <= ZONE_COUNT:
                raise CommonDataError(f"OD line {line_number} has out-of-range TAZ")
            index = (origin - 1) * ZONE_COUNT + (destination - 1)
            if seen[index]:
                raise CommonDataError(f"Duplicate OD pair at line {line_number}: {(origin, destination)}")
            values = [float(value) for value in fields[4:9]]
            if any(value < 0 or not math.isfinite(value) for value in values):
                raise CommonDataError(f"Invalid OD numeric value at line {line_number}")
            seen[index] = 1
            matrix[index] = math.fsum(values)
            for category_index, value in enumerate(values):
                outbound[origin - 1][category_index] += value
            origin_zone_ids.setdefault(origin, fields[1])
            destination_zone_ids.setdefault(destination, fields[3])
            if origin_zone_ids[origin] != fields[1] or destination_zone_ids[destination] != fields[3]:
                raise CommonDataError(f"Inconsistent ZONE_ID for TAZ at line {line_number}")
            record_count += 1
    expected = ZONE_COUNT * ZONE_COUNT
    if record_count != expected or seen.count(1) != expected:
        raise CommonDataError(f"OD coverage failed: records={record_count}, unique={seen.count(1)}, expected={expected}")
    return {
        "matrix": matrix,
        "outbound": outbound,
        "metadata": {
            "record_count": record_count,
            "unique_pair_count": seen.count(1),
            "origin_taz_count": len(origin_zone_ids),
            "destination_taz_count": len(destination_zone_ids),
            "source_sha256": sha256_file(od_path),
            "source_size_bytes": od_path.stat().st_size,
            "source_file": str(od_path.relative_to(ROOT)).replace("\\", "/"),
            "scenario_year": 2023,
            "categories": list(PURPOSES),
            "matrix_value": "sum of five purpose fields",
            "unit": "여객통행/일",
            "orientation": {"records": "origin TAZ to destination TAZ"},
        },
    }


def _only_member(archive: zipfile.ZipFile, token: str) -> str:
    matches = [name for name in archive.namelist() if token in name]
    if len(matches) != 1:
        raise CommonDataError(f"Expected one ZIP member containing {token!r}; found {matches}")
    return matches[0]


def build_national_indicators() -> dict[str, Any]:
    purpose_categories = ("출근", "등교", "업무", "귀가", "기타")
    mode_categories = ("승용차", "버스", "지하철", "일반철도", "고속철도", "항공", "해운")
    with zipfile.ZipFile(NATIONAL_PURPOSE_BUNDLE) as purpose_archive, zipfile.ZipFile(NATIONAL_MODE_BUNDLE) as mode_archive:
        zone_member = _only_member(purpose_archive, "250 존체계.xlsx")
        social_member = _only_member(purpose_archive, "전국지역간_사회경제지표.xlsx")
        purpose_member = _only_member(purpose_archive, "목적별OD(250).xlsx")
        mode_member = _only_member(mode_archive, "주수단별OD(250).xlsx")
        zone_wb = load_workbook(io.BytesIO(purpose_archive.read(zone_member)), read_only=True, data_only=True)
        zone_rows = zone_wb.active.iter_rows(values_only=True)
        if tuple(next(zone_rows)[:5]) != ("대존", "소존", "250존체계", "161존체계", "17존체계"):
            raise CommonDataError("Unexpected national 250 zone header")
        zones = [{"province": row[0], "city_county": row[1], "zone_250": int(row[2]), "zone_161": int(row[3]), "zone_17": int(row[4])} for row in zone_rows]
        zone_wb.close()
        if len(zones) != 250 or {row["zone_250"] for row in zones} != set(range(1, 251)):
            raise CommonDataError("National zone workbook does not contain unique 1..250")

        social_wb = load_workbook(io.BytesIO(purpose_archive.read(social_member)), read_only=True, data_only=True)
        social_rows = social_wb["총인구"].iter_rows(values_only=True)
        social_header = list(next(social_rows))
        year_columns = {year: social_header.index(f"{year}년") for year in NATIONAL_YEARS}
        population = {str(year): {} for year in NATIONAL_YEARS}
        for row in social_rows:
            zone_id = int(row[2])
            for year, column in year_columns.items():
                population[str(year)][str(zone_id)] = float(row[column])
        social_wb.close()
        population_missing = {
            str(year): sorted(set(range(1, 251)) - {int(zone_id) for zone_id in values})
            for year, values in population.items()
        }

        def aggregate(archive: zipfile.ZipFile, member: str, suffix: str, categories: tuple[str, ...]) -> dict[str, Any]:
            workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
            result: dict[str, Any] = {}
            for year in NATIONAL_YEARS:
                sheet = workbook[f"{year}{suffix}"]
                rows = sheet.iter_rows(values_only=True)
                header = list(next(rows))
                if header[:4] != ["출발시도", "도착시도", "출발시군구", "도착시군구"] or header[4:4 + len(categories)] != list(categories):
                    raise CommonDataError(f"Unexpected national indicator header for {year}{suffix}")
                sums = {zone_id: [0.0] * len(categories) for zone_id in range(1, 251)}
                seen = bytearray(250 * 250)
                record_count = 0
                for row in rows:
                    origin, destination = int(row[2]), int(row[3])
                    index = (origin - 1) * 250 + (destination - 1)
                    if seen[index]:
                        raise CommonDataError(f"Duplicate national pair {(origin, destination)} in {year}{suffix}")
                    seen[index] = 1
                    for category_index in range(len(categories)):
                        value = row[4 + category_index]
                        if value is None or isinstance(value, str):
                            raise CommonDataError(f"Non-numeric national value in {year}{suffix}")
                        sums[origin][category_index] += float(value)
                    record_count += 1
                if record_count != 62_500 or seen.count(1) != 62_500:
                    raise CommonDataError(f"National matrix coverage failed for {year}{suffix}")
                result[str(year)] = {
                    str(zone_id): {category: sums[zone_id][i] for i, category in enumerate(categories)}
                    for zone_id in range(1, 251)
                }
            workbook.close()
            return result

        purpose = aggregate(purpose_archive, purpose_member, "_목적OD", purpose_categories)
        mode = aggregate(mode_archive, mode_member, "_주수단OD", mode_categories)
        return {
            "schema_version": 1,
            "zone_system": "national-2024-250",
            "years": list(NATIONAL_YEARS),
            "zones": zones,
            "population": population,
            "population_missing_zone_ids": population_missing,
            "purpose": purpose,
            "main_mode": mode,
            "categories": {"purpose": list(purpose_categories), "main_mode": list(mode_categories)},
            "source": {
                "purpose_bundle": {"file": os.path.relpath(NATIONAL_PURPOSE_BUNDLE, ROOT).replace("\\", "/"), "sha256": sha256_file(NATIONAL_PURPOSE_BUNDLE), "size_bytes": NATIONAL_PURPOSE_BUNDLE.stat().st_size},
                "mode_bundle": {"file": os.path.relpath(NATIONAL_MODE_BUNDLE, ROOT).replace("\\", "/"), "sha256": sha256_file(NATIONAL_MODE_BUNDLE), "size_bytes": NATIONAL_MODE_BUNDLE.stat().st_size},
                "unit": "여객통행/일 (평일 평균 O/D; AAWDT)",
                "orientation": "selected origin zone to all 250 destinations",
            },
        }


def _national_region_target(province: str, city_county: str) -> tuple[str, str, str, dict[str, Any]]:
    """Translate only verified reference-year administrative changes."""
    if province == "강원도":
        return (
            "강원특별자치도",
            city_county,
            "PROVINCE_NAME_CHANGE",
            {
                "effective_date": "2023-06-11",
                "boundary_change": False,
                "evidence": "강원특별자치도청: 종전 강원도가 2023-06-11 강원특별자치도로 출범",
                "url": "https://state.gwd.go.kr/portal/introduce/specialgangwon/introduce",
            },
        )
    if (province, city_county) == ("경상북도", "군위군"):
        return (
            "대구광역시",
            "군위군",
            "AFFILIATION_CHANGE",
            {
                "effective_date": "2023-07-01",
                "boundary_change": True,
                "evidence": "법률 제19155호에 따라 경상북도 군위군이 대구광역시로 편입",
                "url": "https://www.law.go.kr/LSW/lsRvsDocListP.do?chrClsCd=010202&lsId=014378&lsRvsGubun=all",
            },
        )
    if (province, city_county) == ("세종특별자치시", "세종특별자치시"):
        return (
            "세종특별자치시",
            "세종시",
            "SOURCE_LABEL_ALIAS",
            {
                "effective_date": None,
                "boundary_change": False,
                "evidence": "2022 행정경계 sgg=36110의 sggnm '세종시'와 두 배포 존표의 상이한 표시명",
                "url": None,
            },
        )
    return province, city_county, "EXACT_NAME_AND_CODE_FAMILY", {
        "effective_date": None,
        "boundary_change": False,
        "evidence": "두 배포 존표의 시도·시군구 표시명이 정확히 일치",
        "url": None,
    }


def build_region_crosswalk(metro_zones: list[dict[str, Any]], national_zones: list[dict[str, Any]]) -> dict[str, Any]:
    national_by_code: dict[str, list[dict[str, Any]]] = {}
    for row in national_zones:
        code = str(row.get("reference_admin_code5") or "")
        if code:
            national_by_code.setdefault(code, []).append(row)
    metro_regions: dict[tuple[str, str], set[str]] = {}
    for zone in metro_zones:
        metro_regions.setdefault((zone["province"], zone["city_county"]), set()).add(zone["admin_code"][:5])
    rows = []
    unresolved = []
    duplicates = []
    for (province, city_county), admin_codes in sorted(metro_regions.items()):
        target_province, target_city, method, evidence = _national_region_target(province, city_county)
        matches = []
        for code in sorted(admin_codes):
            matches.extend(national_by_code.get(code, []))
        matches = list({int(row["zone_250"]): row for row in matches}.values())
        name_consistent = len(matches) == 1 and (matches[0]["province"], matches[0]["city_county"]) == (target_province, target_city)
        if len(matches) != 1 or not name_consistent:
            item = {
                "metro_region": {"province": province, "city_county": city_county, "admin_code5": sorted(admin_codes)},
                "status": "UNRESOLVED" if not matches or not name_consistent else "DUPLICATE",
                "candidate_count": len(matches),
                "method": method,
                "join_key": "reference_admin_code5",
                "expected_national_name": {"province": target_province, "city_county": target_city},
                "evidence": evidence,
            }
            (unresolved if not matches else duplicates).append(item)
            rows.append(item)
            continue
        match = matches[0]
        rows.append({
            "metro_region": {"province": province, "city_county": city_county, "admin_code5": sorted(admin_codes)},
            "national_region": {"province": match["province"], "city_county": match["city_county"], "zone_250": int(match["zone_250"])},
            "status": "VERIFIED",
            "method": method,
            "join_key": "reference_admin_code5",
            "reference_admin_code5": match["reference_admin_code5"],
            "evidence": evidence,
        })
    zone_ids = [row["national_region"]["zone_250"] for row in rows if row["status"] == "VERIFIED"]
    return {
        "schema_version": 2,
        "source_zone_system": "metro-2024-1310",
        "target_zone_system": "national-2024-250",
        "reference_code_date": "20220101",
        "rows": rows,
        "audit": {
            "source_region_count": len(metro_regions),
            "verified_count": sum(row["status"] == "VERIFIED" for row in rows),
            "unresolved_count": len(unresolved),
            "duplicate_count": len(duplicates),
            "distinct_target_zone_count": len(set(zone_ids)),
            "unresolved": unresolved,
            "duplicates": duplicates,
        },
    }


def build_national_zone_points(national_zones: list[dict[str, Any]], boundary_path: Path = BOUNDARY_REFERENCE_GEOJSON) -> dict[str, Any]:
    features = json.loads(boundary_path.read_text(encoding="utf-8"))["features"]
    by_name: dict[tuple[str, str], list[dict[str, Any]]] = {}
    normalize = lambda value: "".join(str(value or "").split())
    for feature in features:
        properties = feature.get("properties") or {}
        by_name.setdefault(
            (normalize(properties.get("sidonm")), normalize(properties.get("sggnm"))), []
        ).append(feature)
    rows = []
    missing = []
    for zone in national_zones:
        province, city = zone["province"], zone["city_county"]
        lookup_province, lookup_city = province, city
        method = "20220101_EXACT_REGION"
        if province == "강원특별자치도":
            lookup_province, method = "강원도", "20220101_PRE_RENAME_REGION"
        elif (province, city) == ("대구광역시", "군위군"):
            lookup_province, method = "경상북도", "20220101_PRE_AFFILIATION_REGION"
        elif (province, city) == ("세종특별자치시", "세종시"):
            lookup_city, method = "세종시", "20220101_SOURCE_LABEL_REGION"
        matches = by_name.get((normalize(lookup_province), normalize(lookup_city)), [])
        matched_codes = sorted({str((feature.get("properties") or {}).get("sgg") or "") for feature in matches} - {""})
        lon, lat, point_method = _representative_from_features(matches)
        item = {
            **zone,
            "region_name": f"{province} {city}",
            "longitude": lon,
            "latitude": lat,
            "coordinate_method": f"{method}_WHITESPACE_NORMALIZED_{point_method}",
            "boundary_version": "20220101",
            "boundary_feature_count": len(matches),
            "reference_admin_code5": matched_codes[0] if len(matched_codes) == 1 else None,
        }
        rows.append(item)
        if lon is None or lat is None or len(matched_codes) != 1:
            missing.append(int(zone["zone_250"]))
    if missing:
        raise CommonDataError(f"National 250 representative point coverage failed: {missing}")
    return {
        "schema_version": 2,
        "zone_system": "national-2024-250",
        "zone_count": 250,
        "zones": rows,
        "missing_coordinate_zone_ids": missing,
        "coordinate_source": {
            "file": str(boundary_path.relative_to(ROOT)).replace("\\", "/"),
            "sha256": sha256_file(boundary_path),
            "version": "20220101",
            "provenance": "SGIS/행정안전부 기반 vuski/admdongkor 가공본",
            "warning": "전국 250존의 대표 위치 산출용 경계이며 O/D 배포 존표 자체의 공식 경계는 아님",
        },
        "representative_point_policy": "province and city/county labels joined after whitespace normalization; matching boundary features aggregated with polygon area-weighted centroid",
    }


def build_national_access_matrices() -> dict[str, Any]:
    categories = ("출근", "등교", "업무", "귀가", "기타")
    matrices: dict[int, array] = {}
    validation: dict[str, Any] = {}
    with zipfile.ZipFile(NATIONAL_PURPOSE_BUNDLE) as archive:
        member = _only_member(archive, "목적별OD(250).xlsx")
        workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
        for year in NATIONAL_YEARS:
            sheet_name = f"{year}_목적OD"
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = list(next(rows))
            if header[:4] != ["출발시도", "도착시도", "출발시군구", "도착시군구"] or header[4:9] != list(categories):
                raise CommonDataError(f"Unexpected national access O/D header in {sheet_name}")
            matrix = array("d", [0.0]) * (250 * 250)
            seen = bytearray(250 * 250)
            origins, destinations = set(), set()
            for line_number, row in enumerate(rows, start=2):
                origin, destination = int(row[2]), int(row[3])
                if not 1 <= origin <= 250 or not 1 <= destination <= 250:
                    raise CommonDataError(f"National O/D out-of-range zone at {sheet_name}!{line_number}")
                index = (origin - 1) * 250 + destination - 1
                if seen[index]:
                    raise CommonDataError(f"Duplicate national O/D pair {(origin, destination)} in {sheet_name}")
                values = row[4:9]
                if any(value is None or isinstance(value, str) for value in values):
                    raise CommonDataError(f"Non-numeric national O/D value in {sheet_name}!{line_number}")
                seen[index] = 1
                matrix[index] = math.fsum(float(value) for value in values)
                origins.add(origin)
                destinations.add(destination)
            if seen.count(1) != 62_500 or origins != set(range(1, 251)) or destinations != set(range(1, 251)):
                raise CommonDataError(f"National O/D coverage failed in {sheet_name}")
            matrices[year] = matrix
            validation[str(year)] = {
                "sheet": sheet_name,
                "record_count": 62_500,
                "unique_pair_count": 62_500,
                "origin_zone_count": 250,
                "destination_zone_count": 250,
            }
        workbook.close()
    return {
        "matrices": matrices,
        "metadata": {
            "zone_system": "national-2024-250",
            "scenario_years": list(NATIONAL_YEARS),
            "categories": list(categories),
            "matrix_value": "sum of five purpose fields",
            "unit": "여객통행/일 (평일 평균 O/D; AAWDT)",
            "orientation": "row origin national-250 zone to column destination national-250 zone",
            "source_file": os.path.relpath(NATIONAL_PURPOSE_BUNDLE, ROOT).replace("\\", "/"),
            "source_sha256": sha256_file(NATIONAL_PURPOSE_BUNDLE),
            "source_size_bytes": NATIONAL_PURPOSE_BUNDLE.stat().st_size,
            "validation": validation,
        },
    }


def prepare_common_data(*, overwrite: bool = False) -> dict[str, Any]:
    for required in (ZONE_WORKBOOK, BOUNDARY_GEOJSON, BOUNDARY_FALLBACK_GEOJSON, PURPOSE_OD, NATIONAL_PURPOSE_BUNDLE, NATIONAL_MODE_BUNDLE):
        if not required.is_file():
            raise CommonDataError(f"Required source is missing: {required}")
    DERIVED.mkdir(parents=True, exist_ok=True)
    outputs = (
        ZONE_JSON,
        MATRIX_BIN,
        PURPOSE_OUTBOUND_JSON,
        NATIONAL_INDICATOR_JSON,
        NATIONAL_ZONE_JSON,
        REGION_CROSSWALK_JSON,
        MANIFEST_JSON,
        *(national_matrix_path(year) for year in NATIONAL_YEARS),
    )
    if any(path.exists() for path in outputs) and not overwrite:
        raise FileExistsError("Derived common data already exists; pass overwrite=True only for an intentional rebuild")
    zone_payload = build_zone_points()
    purpose = build_purpose_matrix()
    ZONE_JSON.write_text(json.dumps(zone_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with MATRIX_BIN.open("wb") as stream:
        purpose["matrix"].tofile(stream)
    outbound_payload = {
        "schema_version": 1,
        "zone_system": "metro-2024-1310",
        "scenario_year": 2023,
        "categories": list(PURPOSES),
        "rows": [
            {"taz": index + 1, "values": {name: values[i] for i, name in enumerate(PURPOSES)}}
            for index, values in enumerate(purpose["outbound"])
        ],
    }
    PURPOSE_OUTBOUND_JSON.write_text(json.dumps(outbound_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    national_payload = build_national_indicators()
    NATIONAL_INDICATOR_JSON.write_text(json.dumps(national_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    national_zone_payload = build_national_zone_points(national_payload["zones"])
    NATIONAL_ZONE_JSON.write_text(json.dumps(national_zone_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    region_crosswalk = build_region_crosswalk(zone_payload["zones"], national_zone_payload["zones"])
    REGION_CROSSWALK_JSON.write_text(json.dumps(region_crosswalk, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    national_access = build_national_access_matrices()
    national_matrix_manifest = {}
    for year, matrix in national_access["matrices"].items():
        path = national_matrix_path(year)
        with path.open("wb") as stream:
            matrix.tofile(stream)
        national_matrix_manifest[str(year)] = {
            "file": str(path.relative_to(ROOT)).replace("\\", "/"),
            "sha256": sha256_file(path),
            "format": "little-endian/native IEEE754 float64, row-major explicit national-250 zone IDs",
            "expected_bytes": 250 * 250 * struct.calcsize("d"),
        }
    manifest = {
        "schema_version": 2,
        "built_at_utc": datetime.now(timezone.utc).isoformat(),
        "zone_system": "metro-2024-1310",
        "zone_count": ZONE_COUNT,
        "sources": {
            "zone_workbook": {"file": str(ZONE_WORKBOOK.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(ZONE_WORKBOOK), "size_bytes": ZONE_WORKBOOK.stat().st_size},
            "boundary": zone_payload["coordinate_source"],
            "purpose_od": purpose["metadata"],
            "national_purpose_bundle": national_payload["source"]["purpose_bundle"],
            "national_mode_bundle": national_payload["source"]["mode_bundle"],
            "national_boundary_reference": national_zone_payload["coordinate_source"],
        },
        "derived": {
            "zones": {"file": str(ZONE_JSON.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(ZONE_JSON)},
            "matrix": {"file": str(MATRIX_BIN.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(MATRIX_BIN), "format": "little-endian/native IEEE754 float64, row-major explicit TAZ IDs", "expected_bytes": ZONE_COUNT * ZONE_COUNT * struct.calcsize("d")},
            "purpose_outbound": {"file": str(PURPOSE_OUTBOUND_JSON.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(PURPOSE_OUTBOUND_JSON)},
            "national_indicators": {"file": str(NATIONAL_INDICATOR_JSON.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(NATIONAL_INDICATOR_JSON)},
            "national_zones": {"file": str(NATIONAL_ZONE_JSON.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(NATIONAL_ZONE_JSON)},
            "region_crosswalk": {"file": str(REGION_CROSSWALK_JSON.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256_file(REGION_CROSSWALK_JSON)},
            "national_matrices": national_matrix_manifest,
        },
        "checks": {
            "zones": {"count": len(zone_payload["zones"]), "missing_coordinate_count": len(zone_payload["missing_coordinate_taz"])},
            "od": purpose["metadata"],
            "national_region_crosswalk": region_crosswalk["audit"],
            "national_od": national_access["metadata"],
        },
    }
    MANIFEST_JSON.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest
