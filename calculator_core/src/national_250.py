from __future__ import annotations

import hashlib
import io
import math
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


YEARS = (2023, 2025, 2030, 2035, 2040, 2045, 2050)
KIND_SPEC = {
    "purpose": {
        "member_token": "목적별OD(250).xlsx",
        "sheet_suffix": "_목적OD",
        "categories": ("출근", "등교", "업무", "귀가", "기타"),
    },
    "main_mode": {
        "member_token": "주수단별OD(250).xlsx",
        "sheet_suffix": "_주수단OD",
        "categories": ("승용차", "버스", "지하철", "일반철도", "고속철도", "항공", "해운"),
    },
}


class National250Error(ValueError):
    pass


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def _only_member(archive: zipfile.ZipFile, token: str) -> str:
    matches = [name for name in archive.namelist() if token in name]
    if len(matches) != 1:
        raise National250Error(f"Expected exactly one ZIP member containing {token!r}; found {matches}")
    return matches[0]


class National250Reader:
    """Strict reader for the verified KTDB 2024 national 250-zone XLSX bundles."""

    def __init__(self, bundle: Path, kind: str):
        if kind not in KIND_SPEC:
            raise National250Error(f"Unsupported national-250 kind: {kind}")
        self.bundle = Path(bundle)
        self.kind = kind
        if not self.bundle.is_file():
            raise National250Error(f"Bundle not found: {self.bundle}")
        self.archive = zipfile.ZipFile(self.bundle)
        self.zone_member = _only_member(self.archive, "250 존체계.xlsx")
        self.social_member = _only_member(self.archive, "전국지역간_사회경제지표.xlsx")
        self.od_member = _only_member(self.archive, KIND_SPEC[kind]["member_token"])

    def _read_member(self, name: str) -> bytes:
        return self.archive.read(name)

    def source_metadata(self) -> dict[str, Any]:
        od_bytes = self._read_member(self.od_member)
        zone_bytes = self._read_member(self.zone_member)
        social_bytes = self._read_member(self.social_member)
        return {
            "bundle": str(self.bundle),
            "bundle_size": self.bundle.stat().st_size,
            "bundle_sha256": _sha256(self.bundle.read_bytes()),
            "od_member": self.od_member,
            "od_member_size": len(od_bytes),
            "od_member_sha256": _sha256(od_bytes),
            "zone_member": self.zone_member,
            "zone_member_sha256": _sha256(zone_bytes),
            "social_member": self.social_member,
            "social_member_sha256": _sha256(social_bytes),
            "distribution": "2024년 전국 여객O/D 보완갱신",
            "zone_system": "2022-12 administrative districts, national 250 zones",
            "unit": "여객통행/일 (평일 평균 O/D; AAWDT)",
            "scenario_years": list(YEARS),
        }

    def zones(self) -> list[dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(self._read_member(self.zone_member)), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = tuple(next(rows)[:5])
        expected = ("대존", "소존", "250존체계", "161존체계", "17존체계")
        if header != expected:
            raise National250Error(f"Unexpected 250-zone header: {header!r}")
        result = [
            {
                "province": values[0],
                "city_county": values[1],
                "zone_250": int(values[2]),
                "zone_161": int(values[3]),
                "zone_17": int(values[4]),
            }
            for values in rows
        ]
        ids = [row["zone_250"] for row in result]
        if len(result) != 250 or len(set(ids)) != 250 or set(ids) != set(range(1, 251)):
            raise National250Error("250-zone workbook does not contain exactly unique IDs 1..250")
        return result

    def resolve_zone(self, province: str, city_county: str) -> dict[str, Any]:
        matches = [
            row for row in self.zones()
            if row["province"] == province and row["city_county"] == city_county
        ]
        if len(matches) != 1:
            raise National250Error(
                f"Expected one national-250 zone for {province} {city_county}; found {matches}"
            )
        return matches[0]

    def resolve_zones(self, province: str, city_prefix: str) -> list[dict[str, Any]]:
        matches = [
            row for row in self.zones()
            if row["province"] == province and str(row["city_county"]).startswith(city_prefix)
        ]
        if not matches:
            raise National250Error(f"No national-250 zones for {province} prefix {city_prefix!r}")
        return matches

    def population(self, zone_250: int, year: int) -> dict[str, Any]:
        result = self.population_many([zone_250], year)
        return {
            **result,
            "zone_250": int(zone_250),
        }

    def population_many(self, zone_ids: list[int], year: int) -> dict[str, Any]:
        if year not in YEARS:
            raise National250Error(f"Population year {year} is not provided; available={YEARS}")
        wanted = {int(zone_id) for zone_id in zone_ids}
        if not wanted:
            raise National250Error("At least one national-250 zone is required")
        workbook = load_workbook(io.BytesIO(self._read_member(self.social_member)), read_only=True, data_only=True)
        sheet = workbook["총인구"]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        year_label = f"{year}년"
        if header[:3] != ["시도코드", "시군구", "시군구 존체계"] or year_label not in header:
            raise National250Error(f"Unexpected population header: {header!r}")
        year_column = header.index(year_label)
        matches = [row for row in rows if int(row[2]) in wanted]
        matched_ids = {int(row[2]) for row in matches}
        if len(matches) != len(wanted) or matched_ids != wanted:
            raise National250Error(
                f"Population rows for zones {sorted(wanted)} are incomplete: matched={sorted(matched_ids)}"
            )
        return {
            "zone_250_ids": sorted(wanted),
            "year": int(year),
            "members": [
                {
                    "province": row[0],
                    "city_county": row[1],
                    "zone_250": int(row[2]),
                    "population": float(row[year_column]),
                }
                for row in matches
            ],
            "population": math.fsum(float(row[year_column]) for row in matches),
            "unit": "persons",
            "source_sheet": "총인구",
            "source_column": year_label,
        }

    def aggregate_outbound(self, zone_250: int, year: int) -> dict[str, Any]:
        result = self.aggregate_outbound_many([zone_250], year)
        return {
            **result,
            "zone_250": int(zone_250),
        }

    def aggregate_outbound_many(self, zone_ids: list[int], year: int) -> dict[str, Any]:
        if year not in YEARS:
            raise National250Error(f"O/D year {year} is not provided; available={YEARS}")
        wanted = {int(zone_id) for zone_id in zone_ids}
        if not wanted:
            raise National250Error("At least one national-250 origin zone is required")
        spec = KIND_SPEC[self.kind]
        workbook = load_workbook(io.BytesIO(self._read_member(self.od_member)), read_only=True, data_only=True)
        sheet_name = f"{year}{spec['sheet_suffix']}"
        if sheet_name not in workbook.sheetnames:
            raise National250Error(f"Missing O/D sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        expected_prefix = ["출발시도", "도착시도", "출발시군구", "도착시군구"]
        categories = list(spec["categories"])
        if header[:4] != expected_prefix or header[4:4 + len(categories)] != categories:
            raise National250Error(f"Unexpected {self.kind} header in {sheet_name}: {header!r}")

        category_values = {category: [] for category in categories}
        pairs: set[tuple[int, int]] = set()
        duplicate_pairs: list[tuple[int, int]] = []
        origins: set[int] = set()
        destinations: set[int] = set()
        target_record_count = 0
        record_count = 0
        for row in rows:
            origin = int(row[2])
            destination = int(row[3])
            pair = (origin, destination)
            if pair in pairs:
                duplicate_pairs.append(pair)
            pairs.add(pair)
            origins.add(origin)
            destinations.add(destination)
            record_count += 1
            if origin in wanted:
                target_record_count += 1
                for index, category in enumerate(categories, start=4):
                    value = row[index]
                    if value is None or isinstance(value, str):
                        raise National250Error(
                            f"Non-numeric {category} value at {sheet_name} row {record_count + 1}: {value!r}"
                        )
                    category_values[category].append(float(value))

        expected_ids = set(range(1, 251))
        if (
            record_count != 62_500
            or len(pairs) != 62_500
            or duplicate_pairs
            or origins != expected_ids
            or destinations != expected_ids
            or target_record_count != 250 * len(wanted)
        ):
            raise National250Error(
                "O/D matrix contract failed: "
                f"records={record_count}, unique_pairs={len(pairs)}, duplicates={len(duplicate_pairs)}, "
                f"origins={len(origins)}, destinations={len(destinations)}, target_records={target_record_count}"
            )

        totals = {category: math.fsum(values) for category, values in category_values.items()}
        return {
            "kind": self.kind,
            "year": int(year),
            "zone_250_ids": sorted(wanted),
            "direction": "outbound (selected origin zones summed over all 250 destinations)",
            "categories": totals,
            "total": math.fsum(totals.values()),
            "unit": "여객통행/일 (평일 평균 O/D; AAWDT)",
            "source_sheet": sheet_name,
            "validation": {
                "record_count": record_count,
                "unique_pair_count": len(pairs),
                "duplicate_pair_count": len(duplicate_pairs),
                "origin_id_count": len(origins),
                "destination_id_count": len(destinations),
                "target_record_count": target_record_count,
            },
        }
