from __future__ import annotations

import json
import math
import os
import re
import subprocess
import sys
import urllib.parse
import urllib.request
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.core import PilotValidationError, create_copy, extract_target_flows, load_json, sha256_file


ROOT = Path(__file__).resolve().parents[1]
PROFILE_CATALOG = ROOT / "catalog" / "address_profiles.json"
SOURCE_CATALOG = ROOT / "catalog" / "data_sources.json"
DRAFTS = ROOT / "direction_drafts"
RUNTIME_CONFIGS = ROOT / "runtime_configs"
DIRECTIONS = ("north", "west", "south", "east")


def region_group_name(region_name: str | None) -> str:
    """Return a review-sized administrative group without implying a direction."""
    if not region_name:
        return "지역명 미확인"
    parts = str(region_name).split()
    if len(parts) == 1:
        return parts[0]
    if parts[0] == "경기도" and len(parts) >= 3 and parts[1].endswith("시") and parts[2].endswith("구"):
        return " ".join(parts[:3])
    return " ".join(parts[:2])


def build_direction_groups(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for row in rows:
        group_name = region_group_name(row.get("region_name"))
        row["group_id"] = group_name
        group = groups.setdefault(group_name, {
            "group_id": group_name,
            "group_name": group_name,
            "zone_count": 0,
            "nonzero_zone_count": 0,
            "inflow_to_target": 0.0,
            "outflow_from_target": 0.0,
            "direction": None,
            "status": "UNASSIGNED",
        })
        group["zone_count"] += 1
        inflow = float(row["inflow_to_target"])
        outflow = float(row["outflow_from_target"])
        if inflow != 0 or outflow != 0:
            group["nonzero_zone_count"] += 1
        group["inflow_to_target"] = math.fsum((group["inflow_to_target"], inflow))
        group["outflow_from_target"] = math.fsum((group["outflow_from_target"], outflow))
    return sorted(groups.values(), key=lambda item: item["group_name"])


def normalize_address(value: str) -> str:
    text = value.strip().lower()
    text = text.replace("특별시", "").replace("광역시", "").replace("경기도", "경기")
    text = re.sub(r"번지\b", "", text)
    return re.sub(r"[^0-9a-z가-힣]", "", text)


def load_catalogs() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    return load_json(PROFILE_CATALOG)["profiles"], load_json(SOURCE_CATALOG)["sources"]


def public_profile(profile: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in profile.items() if key not in {"aliases", "config_template", "od_draft_source"}}


def find_direction_approval(profile: dict[str, Any], source: dict[str, Any] | None) -> dict[str, Any] | None:
    """Return only an approval still bound to the same target, zone system and source hash."""
    if not source:
        return None
    source_path = _source_path(source)
    for path in sorted(DRAFTS.glob(f"{profile['id']}_direction_approved*.json"), reverse=True):
        try:
            approval = load_json(path)
            same_contract = (
                approval.get("status") == "DIRECTION_APPROVED"
                and int(approval.get("target_taz")) == int(profile["taz"])
                and approval.get("profile", {}).get("zone_system") == profile["zone_system"]
                and approval.get("source", {}).get("id") == source["id"]
                and source_path.is_file()
                and approval.get("source", {}).get("sha256") == sha256_file(source_path)
            )
            if same_contract:
                return {"file": str(path), "status": approval["status"], "approval": approval.get("approval"),
                        "checks": approval.get("checks"), "source_sha256": approval["source"]["sha256"]}
        except (OSError, ValueError, TypeError, KeyError, json.JSONDecodeError):
            continue
    return None


def _juso_search(address: str) -> dict[str, Any] | None:
    key = os.environ.get("JUSO_SEARCH_CONFM_KEY")
    if not key:
        return None
    params = urllib.parse.urlencode({
        "confmKey": key,
        "currentPage": 1,
        "countPerPage": 10,
        "keyword": address,
        "resultType": "json",
    })
    request = urllib.request.Request(
        "https://business.juso.go.kr/addrlink/addrLinkApi.do?" + params,
        headers={"User-Agent": "indicator-od-automation/1.0"},
    )
    with urllib.request.urlopen(request, timeout=15) as response:
        payload = json.loads(response.read().decode("utf-8"))
    common = payload.get("results", {}).get("common", {})
    if str(common.get("errorCode")) != "0":
        raise PilotValidationError(f"공식 주소 API 오류: {common.get('errorMessage', 'unknown error')}")
    rows = payload.get("results", {}).get("juso") or []
    return {
        "provider": "도로명주소 안내시스템 주소검색 API",
        "candidate_count": len(rows),
        "candidates": [
            {
                "road_address": row.get("roadAddr"),
                "parcel_address": row.get("jibunAddr"),
                "building_management_number": row.get("bdMgtSn"),
                "legal_dong_code_admCd": row.get("admCd"),
                "coordinate_fields": {
                    name: row.get(name) for name in ("admCd", "rnMgtSn", "udrtYn", "buldMnnm", "buldSlno")
                },
            }
            for row in rows
        ],
        "coordinate_status": "별도 좌표 API 키와 UTM-K 좌표 변환이 필요함",
    }


def resolve_address(address: str, target_year: int | None = None) -> dict[str, Any]:
    profiles, sources = load_catalogs()
    normalized = normalize_address(address)
    matches = [
        profile for profile in profiles
        if normalized in {normalize_address(alias) for alias in profile["aliases"]}
    ]
    api_result = None
    api_error = None
    try:
        api_result = _juso_search(address)
    except Exception as exc:  # API failure must not silently change the offline result.
        api_error = str(exc)
    if len(matches) != 1:
        return {
            "status": "ADDRESS_CONFIRMATION_REQUIRED" if matches or api_result else "UNSUPPORTED_ADDRESS",
            "input_address": address,
            "normalized_address": normalized,
            "matched_profiles": [public_profile(profile) for profile in matches],
            "official_address_lookup": api_result,
            "official_address_lookup_error": api_error,
            "message": (
                "복수 프로필이 일치하여 사용자가 대상을 선택해야 합니다."
                if len(matches) > 1 else
                "공식 주소 후보는 확인했지만 검증된 TAZ 대응표와 프로젝트 매핑이 없어 자동 확정하지 않습니다."
                if api_result and api_result["candidate_count"] else
                "현재 카탈로그에서 이 주소와 검증된 TAZ·지표·방향 매핑을 찾지 못했습니다."
            ),
        }
    profile = matches[0]
    year_supported = target_year is None or target_year in profile.get("supported_target_years", [])
    data_source = next((item for item in sources if item["id"] == profile.get("od_draft_source")), None)
    direction_approval = find_direction_approval(profile, data_source)
    if profile["workflow_status"] == "READY" and year_supported:
        status = "READY"
    elif not year_supported:
        status = "UNSUPPORTED_TARGET_YEAR"
    else:
        status = profile["workflow_status"]
    return {
        "status": status,
        "input_address": address,
        "normalized_address": normalized,
        "profile": public_profile(profile),
        "available_od": data_source,
        "direction_approval": direction_approval,
        "official_address_lookup": api_result,
        "official_address_lookup_error": api_error,
        "target_year_supported": year_supported,
        "message": "계산 및 Excel 생성 가능" if status == "READY" else "; ".join(profile.get("blocking_reasons", [])) or "지원 조건 확인 필요",
    }


def _unique_path(folder: Path, stem: str, suffix: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    candidate = folder / f"{stem}{suffix}"
    index = 1
    while candidate.exists():
        candidate = folder / f"{stem}_{index}{suffix}"
        index += 1
    return candidate


def build_runtime_config(address: str, target_year: int, project_name: str | None = None,
                         output_name: str | None = None) -> tuple[Path, dict[str, Any]]:
    resolved = resolve_address(address, target_year)
    if resolved["status"] != "READY":
        raise PilotValidationError(f"주소 기반 출력 준비가 완료되지 않았습니다: {resolved['status']} - {resolved['message']}")
    profiles, _ = load_catalogs()
    profile = next(item for item in profiles if item["id"] == resolved["profile"]["id"])
    config = deepcopy(load_json(ROOT / profile["config_template"]))
    safe_id = re.sub(r"[^0-9a-z-]+", "-", profile["id"].lower()).strip("-")
    config["project_id"] = f"address-{safe_id}-{target_year}"
    config["project_name"] = project_name or f"주소입력 시범 {profile['label']} {target_year}"
    config["indicator_target_year"] = int(target_year)
    if output_name:
        output_path = ROOT / "pilot_output" / output_name
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
    else:
        output_path = _unique_path(ROOT / "pilot_output", f"주소입력_{profile['label']}_{target_year}", ".xlsx")
    config["output_workbook"] = str(output_path.relative_to(ROOT)).replace("\\", "/")
    config["address_context"] = {
        "input_address": address,
        "confirmed_address": profile["confirmed_address"],
        "address_status": profile["address_status"],
        "address_scope": profile["address_scope"],
        "address_basis": profile["address_basis"],
        "address_basis_url": profile["address_basis_url"],
        "coordinate_display": "미확인 - 공식 좌표 API 인증정보 없음" if profile.get("coordinates") is None else str(profile["coordinates"]),
        "administrative_dong_code": profile["administrative_dong_code"],
        "legal_dong_code": profile.get("legal_dong_code"),
        "zone_system": profile["zone_system"],
        "taz": profile["taz"],
        "taz_status": profile["taz_status"],
        "taz_basis": profile["taz_basis"],
        "indicator_scope": profile["indicator_scope"],
    }
    config_path = _unique_path(RUNTIME_CONFIGS, config["project_id"], ".json")
    config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return config_path, config


def create_address_output(address: str, target_year: int, project_name: str | None = None,
                          output_name: str | None = None) -> dict[str, Any]:
    config_path, config = build_runtime_config(address, target_year, project_name, output_name)
    result = create_copy(config_path)
    review_path = _unique_path(ROOT / "reports", f"address_output_review_{config['project_id']}", ".json")
    reviewed = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "review_indicator_od_output.py"), "--config", str(config_path),
         "--workbook", result["output_workbook"], "--report", str(review_path)],
        cwd=ROOT, text=True, capture_output=True, encoding="utf-8", errors="replace",
    )
    if reviewed.returncode != 0:
        raise PilotValidationError(
            f"새 Excel은 생성되었지만 독립 review-output을 통과하지 못했습니다. 완료본으로 사용하지 마십시오: {review_path}\n"
            f"{reviewed.stdout[-1000:]}{reviewed.stderr[-1000:]}"
        )
    review = load_json(review_path)
    if review.get("status") != "PASS":
        raise PilotValidationError(f"독립 review-output 상태가 PASS가 아닙니다: {review_path}")
    result["runtime_config"] = str(config_path)
    result["address_context"] = config["address_context"]
    result["independent_review"] = str(review_path)
    result["independent_review_markdown"] = str(review_path.with_suffix(".md"))
    result["independent_review_status"] = review["status"]
    return result


def _source_path(source: dict[str, Any]) -> Path:
    return (ROOT / source["file"]).resolve()


def generate_direction_draft(address: str) -> dict[str, Any]:
    resolved = resolve_address(address)
    profile = resolved.get("profile")
    source = resolved.get("available_od")
    if not profile or not source:
        raise PilotValidationError("이 주소에는 로컬에서 읽을 수 있는 새 대상 O/D 원자료가 등록되어 있지 않습니다.")
    if source["availability"] not in {"LOCAL_EXTRACTION_ONLY", "LOCAL_AND_VERIFIED"}:
        raise PilotValidationError(f"O/D 원자료 상태가 추출 가능하지 않습니다: {source['availability']}")
    source_path = _source_path(source)
    if not source_path.is_file():
        raise PilotValidationError(f"등록된 O/D 원자료가 로컬에 없습니다: {source_path}")
    zone_path = ROOT / "inputs" / "260608_230912_미아중심 지표.xlsx"
    zone_wb = load_workbook(zone_path, read_only=True, data_only=True)
    zone_ws = zone_wb["존체계"]
    zones: dict[int, dict[str, Any]] = {}
    for row in range(1, zone_ws.max_row + 1):
        raw_taz = zone_ws.cell(row, 6).value
        if isinstance(raw_taz, (int, float)) and not isinstance(raw_taz, bool):
            zones[int(raw_taz)] = {"name": zone_ws.cell(row, 2).value, "administrative_code": zone_ws.cell(row, 7).value}
    zone_wb.close()
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook[source["sheet"]]
    value_rows = sheet.iter_rows(values_only=True)
    header = next(value_rows)
    destination_ids = [int(value) for value in header[1:] if isinstance(value, (int, float)) and not isinstance(value, bool)]
    if len(destination_ids) != len(set(destination_ids)):
        workbook.close()
        raise PilotValidationError("O/D 도착 TAZ 헤더에 중복이 있습니다.")
    target_taz = int(profile["taz"])
    if target_taz not in destination_ids:
        workbook.close()
        raise PilotValidationError(f"대상 TAZ {target_taz}가 O/D 도착 헤더에 없습니다.")
    target_column_index = destination_ids.index(target_taz) + 1
    inflow: dict[int, float] = {}
    target_row_values = None
    for values in value_rows:
        origin_raw = values[0]
        if not isinstance(origin_raw, (int, float)) or isinstance(origin_raw, bool):
            workbook.close()
            raise PilotValidationError(f"O/D 출발 TAZ가 숫자가 아닙니다: {origin_raw!r}")
        origin = int(origin_raw)
        if origin in inflow:
            workbook.close()
            raise PilotValidationError(f"O/D 출발 TAZ {origin}이 중복되었습니다.")
        inflow[origin] = float(values[target_column_index])
        if origin == target_taz:
            target_row_values = values
    workbook.close()
    if target_row_values is None:
        raise PilotValidationError(f"대상 TAZ {target_taz}가 O/D 출발 행에 없습니다.")
    outflow = {taz: float(target_row_values[index + 1]) for index, taz in enumerate(destination_ids)}
    row_map = {taz: index for index, taz in enumerate(inflow)}
    column_map = {taz: index for index, taz in enumerate(destination_ids)}
    if len(row_map) != source["zone_count"] or len(column_map) != source["zone_count"] or set(row_map) != set(column_map):
        raise PilotValidationError("O/D 행·열 TAZ 집합 또는 존 수가 카탈로그와 다릅니다.")
    rows = []
    for taz in sorted(row_map):
        zone = zones.get(taz, {})
        rows.append({
            "taz": taz,
            "region_name": zone.get("name"),
            "administrative_code": zone.get("administrative_code"),
            "inflow_to_target": inflow[taz],
            "outflow_from_target": outflow[taz],
            "direction": "internal" if taz == int(profile["taz"]) else None,
            "status": "INTERNAL_EXCLUDED" if taz == int(profile["taz"]) else "UNASSIGNED",
        })
    payload = {
        "schema_version": 1,
        "status": "DIRECTION_CONFIRMATION_REQUIRED",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "input_address": address,
        "profile": profile,
        "source": {**source, "resolved_file": str(source_path), "sha256": sha256_file(source_path), "size_bytes": source_path.stat().st_size},
        "orientation": source["orientation"],
        "target_taz": int(profile["taz"]),
        "internal_trip": inflow[int(profile["taz"])],
        "rows": rows,
        "groups": build_direction_groups(rows),
        "assignment_basis": {
            "status": "NOT_SELECTED",
            "approved_method": None,
            "required_choice": [
                {
                    "id": "geographic_bearing",
                    "label": "사업지 기준 지리적 방위",
                    "requirement": "사업지 대표점과 각 TAZ의 검증된 대표 좌표 또는 경계가 필요하며, 자동 배정은 승인 전 초안으로만 사용",
                },
                {
                    "id": "access_road_cordon",
                    "label": "주요 접근도로·코든 기준",
                    "requirement": "북·서·남·동으로 볼 접근 도로축과 각 지역/TAZ가 어느 도로축으로 유입·유출하는지에 대한 업무 기준이 필요",
                },
            ],
        },
        "checks": {
            "origin_count": len(row_map),
            "destination_count": len(column_map),
            "sets_equal": set(row_map) == set(column_map),
            "unassigned_count": len(row_map) - 1,
            "assigned_count": 0,
            "excluded_count": 1,
        },
        "warning": "방위 또는 도로 접근 기준을 자동 추정하지 않았습니다. 행정구역 묶음은 검토 편의를 위한 그룹일 뿐 방향 근거가 아니며, 기준 승인 뒤 모든 비내부 존을 네 방향 또는 명시적 제외로 확인해야 접근강도를 계산할 수 있습니다.",
    }
    path = _unique_path(DRAFTS, f"{profile['id']}_direction_draft", ".json")
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"status": payload["status"], "draft": str(path), "checks": payload["checks"], "target_taz": payload["target_taz"],
            "source": payload["source"], "warning": payload["warning"]}


def save_direction_assignment(draft_path: Path, assignments: dict[str, str], approval_note: str) -> dict[str, Any]:
    payload = load_json(draft_path)
    if not approval_note.strip():
        raise PilotValidationError("방향 배정의 적용 기준 또는 승인 메모가 필요합니다.")
    allowed = {*DIRECTIONS, "excluded"}
    by_taz = {str(row["taz"]): row for row in payload["rows"]}
    unknown = sorted(set(assignments) - set(by_taz))
    if unknown:
        raise PilotValidationError(f"원자료에 없는 TAZ가 배정에 포함되었습니다: {unknown[:10]}")
    for taz, direction in assignments.items():
        if int(taz) == int(payload["target_taz"]):
            if direction != "internal":
                raise PilotValidationError("대상 내부 TAZ는 internal로 유지해야 합니다.")
            continue
        if direction not in allowed:
            raise PilotValidationError(f"TAZ {taz}의 방향이 잘못되었습니다: {direction!r}")
        by_taz[taz]["direction"] = direction
        by_taz[taz]["status"] = "EXCLUDED_CONFIRMED" if direction == "excluded" else "ASSIGNED"
    unassigned = [row["taz"] for row in payload["rows"] if row["taz"] != payload["target_taz"] and row.get("direction") not in allowed]
    if unassigned:
        raise PilotValidationError(f"미배정 존 {len(unassigned)}개가 남아 있습니다. 예: {unassigned[:10]}")
    totals = {direction: {"inflow": 0.0, "outflow": 0.0, "count": 0} for direction in DIRECTIONS}
    excluded = []
    for row in payload["rows"]:
        direction = row["direction"]
        if direction in DIRECTIONS:
            totals[direction]["inflow"] = math.fsum((totals[direction]["inflow"], float(row["inflow_to_target"])))
            totals[direction]["outflow"] = math.fsum((totals[direction]["outflow"], float(row["outflow_from_target"])))
            totals[direction]["count"] += 1
        elif direction in {"internal", "excluded"}:
            excluded.append(row["taz"])
    denominator_in = math.fsum(item["inflow"] for item in totals.values())
    denominator_out = math.fsum(item["outflow"] for item in totals.values())
    if denominator_in == 0 or denominator_out == 0:
        raise PilotValidationError("배정 결과의 유입 또는 유출 분모가 0입니다.")
    for item in totals.values():
        item["share_inflow"] = item["inflow"] / denominator_in
        item["share_outflow"] = item["outflow"] / denominator_out
    payload["status"] = "DIRECTION_APPROVED"
    payload["approval"] = {"approved_at": datetime.now().astimezone().isoformat(timespec="seconds"), "note": approval_note}
    payload["direction_totals"] = totals
    payload["denominators"] = {"inflow": denominator_in, "outflow": denominator_out}
    payload["excluded_taz"] = sorted(excluded)
    payload["checks"].update({
        "unassigned_count": 0,
        "assigned_count": sum(item["count"] for item in totals.values()),
        "excluded_count": len(excluded),
        "share_sum_inflow": math.fsum(item["share_inflow"] for item in totals.values()),
        "share_sum_outflow": math.fsum(item["share_outflow"] for item in totals.values()),
    })
    approval_path = _unique_path(DRAFTS, f"{payload['profile']['id']}_direction_approved", ".json")
    approval_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"status": payload["status"], "approval": str(approval_path), "direction_totals": totals,
            "denominators": payload["denominators"], "checks": payload["checks"]}
